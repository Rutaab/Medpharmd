import re
import sys
import io
import tempfile
from pathlib import Path
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

try:
    from PIL import Image
    import pytesseract
    HAS_OCR = True
except ImportError:
    HAS_OCR = False


def clean_ocr(text):
    result = (text
        .replace('|', 'I')
        .replace('٠', 'o')
        .replace('\u2013', '-')
        .replace('\u2014', '--')
        .replace('\u2018', "'")
        .replace('\u2019', "'")
        .replace('\u201c', '"')
        .replace('\u201d', '"')
    )
    result = re.sub(r'^Il\.', 'II.', result)
    result = re.sub(r'^lll\.', 'III.', result)
    result = re.sub(r'^iv\.', 'IV.', result)
    result = re.sub(r'^vl\.', 'VI.', result)
    result = re.sub(r'^xl\.', 'XL.', result)
    result = re.sub(r'^lx\.', 'LX.', result)
    result = re.sub(r'^xc\.', 'XC.', result)
    result = re.sub(r'\bIl\.\s+[A-Z]', 'II. ', result)
    result = re.sub(r'\blll\.\s+[A-Z]', 'III. ', result)
    result = re.sub(r'\b0-([A-Z][a-z])', r'O-\1', result)
    return result


def heading_level(style_name):
    if not style_name:
        return None
    name = style_name.strip().lower()
    if name in ('heading 1', 'heading1', 'title'):
        return 1
    if name in ('heading 2', 'heading2', 'subtitle'):
        return 2
    if name in ('heading 3', 'heading3'):
        return 3
    if name in ('heading 4', 'heading4'):
        return 4
    if name.startswith('heading'):
        try:
            return int(name.split()[-1])
        except (ValueError, IndexError):
            return None
    return None


def run_text(run):
    text = run.text
    if run.bold and run.italic:
        text = f'***{text}***'
    elif run.bold:
        text = f'**{text}**'
    elif run.italic:
        text = f'*{text}*'
    if run.font.superscript:
        text = f'^{text}^'
    if run.font.subscript:
        text = f'~{text}~'
    return text


def paragraph_to_md(para):
    parts = [run_text(run) for run in para.runs]
    return ''.join(parts)


def list_marker(para, index, list_state):
    style = para.style.name.strip().lower() if para.style and para.style.name else ''
    if 'list bullet' in style or 'unordered' in style or 'bullet' in style:
        return '- '
    if 'list number' in style or 'ordered' in style or 'enum' in style:
        return f'{index}. '
    if para.text.strip().startswith('- ') or para.text.strip().startswith('* ') or para.text.strip().startswith('+ '):
        return '- '
    m = re.match(r'^(\d+)[.)]\s', para.text.strip())
    if m:
        return f'{m.group(1)}. '
    return None


def table_to_md(table):
    rows = []
    for row_idx, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            text = cell.text.strip().replace('\n', ' ')
            cells.append(text)
        rows.append(cells)

    if not rows:
        return ''

    col_count = max(len(r) for r in rows)
    header = rows[0]
    while len(header) < col_count:
        header.append('')

    lines = []
    lines.append('| ' + ' | '.join(header) + ' |')
    lines.append('| ' + ' | '.join(['---'] * col_count) + ' |')

    for row in rows[1:]:
        while len(row) < col_count:
            row.append('')
        lines.append('| ' + ' | '.join(row) + ' |')

    return '\n'.join(lines)


def extract_images_ocr(doc):
    """Extract images from docx and run OCR on them."""
    if not HAS_OCR:
        return []
    texts = []
    for rel_id, rel in doc.part.rels.items():
        if 'image' in rel.reltype:
            try:
                image_data = rel.target_part.blob
                img = Image.open(io.BytesIO(image_data))
                text = pytesseract.image_to_string(img)
                if text.strip():
                    texts.append(clean_ocr(text.strip()))
            except Exception as e:
                texts.append(f'[OCR Error: {e}]')
    return texts


def convert_docx_to_md(input_path, output_path=None, enable_ocr=True):
    input_path = Path(input_path)
    if not input_path.exists():
        print(f"Error: File not found: {input_path}")
        return

    if output_path is None:
        output_path = input_path.with_suffix('.md')

    doc = Document(str(input_path))
    md_lines = []
    list_counter = 0
    in_list = False
    in_table = False

    # Process paragraphs
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        style_name = para.style.name if para.style and para.style.name else ''
        hl = heading_level(style_name)

        if in_table:
            in_table = False
            md_lines.append('')

        if hl:
            in_list = False
            list_counter = 0
            prefix = '#' * hl
            md_lines.append(f'{prefix} {clean_ocr(text)}')
            md_lines.append('')
        elif not text:
            in_list = False
            if md_lines and md_lines[-1] != '':
                md_lines.append('')
        else:
            marker = list_marker(para, list_counter + 1 if in_list else 1, None)
            if marker is not None:
                if not in_list:
                    in_list = True
                    list_counter = 1
                else:
                    list_counter += 1
                md_text = paragraph_to_md(para)
                md_text = re.sub(r'^[-*+]\s|\d+[.)]\s', '', md_text, count=1)
                md_lines.append(f'{marker}{clean_ocr(md_text)}')
            else:
                in_list = False
                md_text = paragraph_to_md(para)
                md_lines.append(clean_ocr(md_text))
                md_lines.append('')

    # Process tables
    for table_idx, table in enumerate(doc.tables):
        if md_lines and md_lines[-1] != '':
            md_lines.append('')
        md_lines.append(table_to_md(table))
        md_lines.append('')

    # OCR for embedded images
    if enable_ocr:
        ocr_texts = extract_images_ocr(doc)
        if ocr_texts:
            if md_lines and md_lines[-1] != '':
                md_lines.append('')
            md_lines.append('---')
            md_lines.append('')
            md_lines.append('## OCR Extracted Content')
            md_lines.append('')
            for t in ocr_texts:
                md_lines.append(t)
                md_lines.append('')

    # Clean up blank lines
    cleaned = []
    blank_count = 0
    for line in md_lines:
        if line == '':
            blank_count += 1
            if blank_count <= 2:
                cleaned.append(line)
        else:
            blank_count = 0
            cleaned.append(line)

    output = '\n'.join(cleaned).strip() + '\n'
    output_path = Path(output_path)
    output_path.write_text(output, encoding='utf-8')
    print(f"Converted: {input_path} -> {output_path}")
    print(f"  Output size: {len(output.splitlines())} lines")


if __name__ == '__main__':
    args = sys.argv[1:]
    if not args:
        print("Usage:")
        print("  python docx_to_md.py <file.docx>              Convert a single .docx file")
        print("  python docx_to_md.py <file.docx> <output.md>  Convert to specific output path")
        print("  python docx_to_md.py <directory>              Convert all .docx in directory")
        print("  python docx_to_md.py <file.docx> --no-ocr     Disable OCR for images")
        sys.exit(1)

    enable_ocr = True
    input_arg = args[0]
    output_arg = None
    remaining = args[1:]

    if '--no-ocr' in remaining:
        enable_ocr = False
        remaining.remove('--no-ocr')

    if remaining:
        output_arg = remaining[0]

    input_path = Path(input_arg)

    if input_path.is_dir():
        for f in sorted(input_path.glob('*.docx')):
            convert_docx_to_md(f, enable_ocr=enable_ocr)
    else:
        convert_docx_to_md(input_path, output_arg, enable_ocr=enable_ocr)
